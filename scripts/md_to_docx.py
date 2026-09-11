"""마크다운 요구사항 문서를 docx 로 변환한다.

docs/*.md 가 원본이고 docx 는 배포본이다. md 를 고친 뒤 이걸 돌려 docx 를 맞춘다.
둘이 어긋나 있던 적이 있어 손으로 다시 만들지 않도록 스크립트로 둔다.

    cd backend
    poetry run python ../scripts/md_to_docx.py

특정 파일만:
    poetry run python ../scripts/md_to_docx.py ../docs/newsflow-srs-v10.md

다루는 문법은 이 문서들이 실제로 쓰는 것만이다 — 제목, 문단, 표, 목록,
코드 블록, 인용, 굵게, 인라인 코드, 수평선.
"""

from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

DOCS = Path(__file__).resolve().parents[1] / "docs"

# **굵게** 와 `코드` 만 다룬다. 문서에서 쓰는 인라인 서식이 이 둘뿐이다.
INLINE = re.compile(r"(\*\*.+?\*\*|`[^`]+`)")


def add_inline(paragraph, text: str) -> None:
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
            run.font.color.rgb = RGBColor(0xC7, 0x25, 0x4E)
        else:
            paragraph.add_run(part)


def is_table_separator(line: str) -> bool:
    return bool(re.fullmatch(r"\s*\|[\s:|-]+\|\s*", line))


def split_row(line: str) -> list[str]:
    return [c.strip() for c in line.strip().strip("|").split("|")]


def add_table(doc: Document, rows: list[list[str]]) -> None:
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    for r, cells in enumerate(rows):
        for c, text in enumerate(cells):
            if c >= len(rows[0]):
                continue
            cell = table.cell(r, c)
            cell.text = ""
            para = cell.paragraphs[0]
            add_inline(para, text)
            for run in para.runs:
                run.font.size = Pt(9)
                if r == 0:
                    run.bold = True
    doc.add_paragraph()


def add_code(doc: Document, lines: list[str]) -> None:
    para = doc.add_paragraph()
    para.paragraph_format.left_indent = Pt(18)
    run = para.add_run("\n".join(lines))
    run.font.name = "Consolas"
    run.font.size = Pt(9)


def convert(md_path: Path, docx_path: Path) -> None:
    doc = Document()
    doc.styles["Normal"].font.name = "맑은 고딕"
    doc.styles["Normal"].font.size = Pt(10)

    lines = md_path.read_text(encoding="utf-8").splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        # 코드 블록
        if stripped.startswith("```"):
            i += 1
            block = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                block.append(lines[i])
                i += 1
            add_code(doc, block)
            i += 1
            continue

        # 표
        if stripped.startswith("|") and i + 1 < len(lines) and is_table_separator(lines[i + 1]):
            header = split_row(stripped)
            i += 2
            rows = [header]
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = split_row(lines[i])
                cells += [""] * (len(header) - len(cells))
                rows.append(cells[: len(header)])
                i += 1
            add_table(doc, rows)
            continue

        # 제목
        heading = re.match(r"(#{1,6})\s+(.*)", stripped)
        if heading:
            level = len(heading.group(1))
            para = doc.add_heading("", level=min(level, 4))
            add_inline(para, heading.group(2))
            i += 1
            continue

        # 수평선
        if stripped in {"---", "***", "___"}:
            para = doc.add_paragraph()
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            para.add_run("─" * 40).font.color.rgb = RGBColor(0xBB, 0xBB, 0xBB)
            i += 1
            continue

        # 인용
        if stripped.startswith(">"):
            para = doc.add_paragraph(style="Intense Quote")
            add_inline(para, stripped.lstrip("> ").strip())
            i += 1
            continue

        # 목록
        bullet = re.match(r"[-*]\s+(.*)", stripped)
        if bullet:
            add_inline(doc.add_paragraph(style="List Bullet"), bullet.group(1))
            i += 1
            continue
        numbered = re.match(r"\d+\.\s+(.*)", stripped)
        if numbered:
            add_inline(doc.add_paragraph(style="List Number"), numbered.group(1))
            i += 1
            continue

        # 빈 줄 / 문단
        if not stripped:
            i += 1
            continue
        add_inline(doc.add_paragraph(), stripped)
        i += 1

    doc.save(docx_path)
    print(f"  {md_path.name} -> {docx_path.name}")


def main(argv: list[str]) -> int:
    targets = [Path(a).resolve() for a in argv[1:]] or sorted(DOCS.glob("newsflow-*.md"))
    if not targets:
        print("변환할 문서가 없다", file=sys.stderr)
        return 1
    print("docx 생성")
    for md in targets:
        convert(md, md.with_suffix(".docx"))
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv))
